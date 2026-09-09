import openpyxl

FOLDER = r'c:\Users\chuch\.gemini\antigravity\playground\compromisosv2'
wb = openpyxl.load_workbook(FOLDER + r'\Pruebareporteejecutivo_validado_PCOM_CM_correcto.xlsx', read_only=True)

print('Hojas:', wb.sheetnames)

ws_bd = wb['Base de datos origen']
bd = list(ws_bd.iter_rows(values_only=True))
headers = bd[1]

# Extraer top contratos con CM
cm_contratos = []
sin_cm = []
for r in bd[2:]:
    if r[0] is None: continue
    no_cont = str(r[5]).strip() if r[5] else ''
    conv_mod = str(r[9]).strip().upper() if r[9] else 'NO'
    proveedor = str(r[7]).strip() if r[7] else 'N/A'
    monto_max = float(r[24]) if isinstance(r[24],(int,float)) else 0.0
    monto_pago = float(r[39]) if isinstance(r[39],(int,float)) else 0.0
    monto_est = float(r[46]) if isinstance(r[46],(int,float)) else 0.0
    ptda = str(r[17]).strip() if r[17] else 'N/A'
    anexo_raw = str(r[6]).strip().split('\n')[0] if r[6] else 'Sin Anexo'
    
    item = {'contrato':no_cont, 'proveedor':proveedor[:50], 'monto_max':monto_max,
            'pagado':monto_pago, 'por_ejercer':monto_est, 'ptda':ptda, 'anexo':anexo_raw}
    if conv_mod in ('SI', 'S'):
        cm_contratos.append(item)
    else:
        sin_cm.append(item)

print('Contratos con CM:', len(cm_contratos))
print('Contratos sin CM:', len(sin_cm))

cm_top = sorted(cm_contratos, key=lambda x: -x['monto_max'])[:8]
print()
print('TOP CONTRATOS CON CONVENIO MODIFICATORIO:')
for c in cm_top:
    print('  ' + c['contrato'][:30] + ' | ' + c['proveedor'][:40] + ' | ' + str(round(c['monto_max'],0)))

# Datos por Anexo detallados
print()
print('=== ESTADISTICAS POR ANEXO (de reporte validado) ===')
anexos_data = {}
for r in bd[2:]:
    if r[0] is None: continue
    anexo = str(r[6]).strip().split('\n')[0] if r[6] else 'Sin Anexo'
    if not anexo: anexo = 'Sin Anexo'
    monto_max = float(r[24]) if isinstance(r[24],(int,float)) else 0.0
    monto_sicop = float(r[35]) if isinstance(r[35],(int,float)) else 0.0
    monto_pago = float(r[39]) if isinstance(r[39],(int,float)) else 0.0
    monto_est = float(r[46]) if isinstance(r[46],(int,float)) else 0.0
    
    if anexo not in anexos_data:
        anexos_data[anexo] = {'regs':0,'techo':0,'sicop':0,'pagado':0,'est':0}
    anexos_data[anexo]['regs'] += 1
    anexos_data[anexo]['techo'] += monto_max
    anexos_data[anexo]['sicop'] += monto_sicop
    anexos_data[anexo]['pagado'] += monto_pago
    anexos_data[anexo]['est'] += monto_est

for k, v in sorted(anexos_data.items(), key=lambda x: -x[1]['techo'])[:12]:
    print('  ' + k + ': regs=' + str(v['regs']) + ' techo=' + str(round(v['techo'],0)) + ' sicop=' + str(round(v['sicop'],0)) + ' pagado=' + str(round(v['pagado'],0)) + ' est=' + str(round(v['est'],0)))
